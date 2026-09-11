#!/usr/bin/env python3
"""Reconciliation twin of the VBA macro RunReconciliation.

Reads ../data/synthetic/ledgers.csv (files A and B) and produces
../workbooks/output/reconciliation_report.xlsx:
  - Exceptions tab: tolerance-banded differences (green = match, amber = within
    tolerance < $5,000, red = material >= $5,000 or missing on one side)
  - Summary tab: counts + totals by band

The banding logic mirrors the VBA macro in vba/RunReconciliation.bas one-for-one,
so a reviewer can diff the two implementations.

Usage:  python3 recon.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "synthetic" / "ledgers.csv"
OUT = ROOT / "workbooks" / "output" / "reconciliation_report.xlsx"
TOLERANCE = 5_000

GOLD = "A68A5B"
GREEN_FILL = PatternFill("solid", fgColor="E3EEE7")
AMBER_FILL = PatternFill("solid", fgColor="F6EEDD")
RED_FILL = PatternFill("solid", fgColor="F3E1E4")
HEAD_FONT = Font(name="Georgia", size=11, bold=True, color="1E4B38")
THIN = Border(bottom=Side(style="thin", color=GOLD))


def band(diff: float) -> str:
    a = abs(diff)
    if a < 0.005:
        return "MATCH"
    if a < TOLERANCE:
        return "WITHIN TOLERANCE"
    return "MATERIAL"


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    df = pd.read_csv(SRC)
    a = df[df.file == "A"][["fund_control_point", "object_class", "description", "amount"]].rename(columns={"amount": "a_amt"})
    b = df[df.file == "B"][["fund_control_point", "object_class", "amount"]].rename(columns={"amount": "b_amt"})

    j = a.merge(b, on=["fund_control_point", "object_class"], how="outer", indicator=True)
    j["diff"] = j["a_amt"] - j["b_amt"]
    j["status"] = j["_merge"].map({"both": "ok", "left_only": "MISSING IN B", "right_only": "MISSING IN A"})
    j["diff"] = j["diff"].where(j["_merge"] == "both", None)
    j["band"] = j["diff"].apply(lambda d: "MISSING" if d is None or pd.isna(d) else band(d))

    j = j.sort_values(["band", "diff"], ascending=[True, False], na_position="last")

    wb = Workbook()
    ws = wb.active
    ws.title = "Exceptions"
    cols = ["fund_control_point", "object_class", "description", "a_amt", "b_amt", "diff", "status", "band"]
    ws.append(["Fund Control Point", "Object Class", "Description", "Ledger A", "Ledger B",
               "Difference (A-B)", "Status", "Band"])
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD_FONT
        cell.border = THIN
    fills = {"MATCH": GREEN_FILL, "WITHIN TOLERANCE": AMBER_FILL, "MATERIAL": RED_FILL, "MISSING": RED_FILL}
    for _, r in j.iterrows():
        ws.append([r["fund_control_point"], r["object_class"], r["description"],
                   None if pd.isna(r["a_amt"]) else r["a_amt"], None if pd.isna(r["b_amt"]) else r["b_amt"],
                   None if pd.isna(r["diff"]) else r["diff"], r["status"], r["band"]])
        row = ws.max_row
        f = fills.get(r["band"])
        if f:
            for c in range(1, len(cols) + 1):
                ws.cell(row=row, column=c).fill = f
        for c in (4, 5, 6):
            ws.cell(row=row, column=c).number_format = "#,##0"
    for i, w in enumerate([20, 14, 34, 14, 14, 16, 14, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    s = wb.create_sheet("Summary")
    s.append(["Reconciliation Summary — synthetic ledgers (seed 20260911)"])
    s["A1"].font = Font(name="Georgia", size=14, bold=True, color="1E4B38")
    s.append([f"Tolerance band: ${TOLERANCE:,}"])
    s.append([])
    s.append(["Band", "Count", "Total |Difference|"])
    for c in range(1, 4):
        s.cell(row=4, column=c).font = HEAD_FONT
    for bname in ["MATCH", "WITHIN TOLERANCE", "MATERIAL", "MISSING"]:
        sub = j[j.band == bname]
        s.append([bname, len(sub), float(sub["diff"].abs().sum()) if len(sub) else 0.0])
    s.append(["TOTAL", len(j), float(j["diff"].abs().sum())])
    for row in s.iter_rows(min_row=5, max_col=3):
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"
    s.column_dimensions["A"].width = 22
    s.column_dimensions["B"].width = 12
    s.column_dimensions["C"].width = 20

    wb.save(OUT)
    print(f"OK -> {OUT} ({len(j)} exception rows; {len(j[j.band != 'MATCH'])} non-matches)")


if __name__ == "__main__":
    main()
